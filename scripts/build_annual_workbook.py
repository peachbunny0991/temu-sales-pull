# -*- coding: utf-8 -*-
"""
Temu 商家中心 → 年度每日销量台账工作簿生成器（可复用）

用法:
    python3 build_annual_workbook.py --store golf|towel [--year 2026] [--out-dir <目录>] [--template <模板xlsx>]

行为:
    1. 扫描 <data_dir>/<store>/ 下的 raw_<YYYYMM>.json + daily_<YYYYMM>.json（逐日数据）
    2. 从空白模板重建年度工作簿：
       - 每个月一个 sheet（"2026-09" 格式），有数据的月份填数据+产品图，无数据的月份为空白模板
       - 月份 sheet 从「最早有数据的月份」到 12 月全部创建
    3. 逐日数据只填「当月已有数据的天」对应列（L列=1日 …），无数据天留空
    4. 输出 <out_dir>/<店名> 每日销量<年>.xlsx

数据约定:
    raw_<YYYYMM>.json   : [{skcId(String), name, skus:[{className, skuId(String), today, d7, d30, total}]}]
    daily_<YYYYMM>.json : {skcId: [{date:"YYYY-MM-DD", prodSkuId, salesNumber, isPredict, soldOut}]}
    imgs/<store>_<skcId>.jpg : 产品主图（800x800），由 download_images.py 生成
"""
import argparse, glob, json, os, re, sys
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment

STORES = {
    "golf":  {"name": "Golf Sports Factory Shop"},
    "towel": {"name": "Towel Manufacturer"},
}

DEFAULT_TPL = os.environ.get("TEMU_TPL")
DEFAULT_OUT = os.environ.get("TEMU_OUT_DIR")
DEFAULT_DATA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data")

def excel_serial_1st(year, month):
    """当月1日的Excel序列号（1899-12-30为0）"""
    from datetime import date
    return (date(year, month, 1) - date(1899, 12, 30)).days

def build(store_key, year, template, out_dir, data_dir):
    if not template or not os.path.isfile(template):
        sys.exit("模板不存在，请传 --template 或设环境变量 TEMU_TPL（指向「2026每日销量模板（空白）.xlsx」）: %r" % template)
    if not out_dir:
        sys.exit("未指定输出目录，请传 --out-dir 或设环境变量 TEMU_OUT_DIR")
    cfg = STORES[store_key]
    store_dir = os.path.join(data_dir, store_key)
    if not os.path.isdir(store_dir):
        sys.exit("数据目录不存在: %s" % store_dir)

    raw_files = sorted(glob.glob(os.path.join(store_dir, "raw_%d*.json" % year)))
    daily_files = sorted(glob.glob(os.path.join(store_dir, "daily_%d*.json" % year)))
    months = sorted({int(re.search(r"(\d{4})(\d{2})\.json$", os.path.basename(f)).group(2)) for f in raw_files})
    if not months:
        sys.exit("未找到 %d 年的数据文件" % year)
    print("店铺:", cfg["name"], "| 有数据的月份:", months)

    # 数据装载: month -> (raw, daily_map)
    data = {}
    for m in months:
        raw = json.load(open(os.path.join(store_dir, "raw_%d%02d.json" % (year, m)), encoding="utf-8"))
        daily = json.load(open(os.path.join(store_dir, "daily_%d%02d.json" % (year, m)), encoding="utf-8"))
        daily_map = {}
        for skc, arr in daily.items():
            mm = {}
            for x in arr:
                mm.setdefault(str(x['prodSkuId']), {})[x['date']] = x['salesNumber']
            daily_map[skc] = mm
        data[m] = (raw, daily_map)

    wb = load_workbook(template)
    tpl = wb.active

    for m in range(min(months), 13):
        ws = wb.copy_worksheet(tpl)
        ws.title = "%d-%02d" % (year, m)
        ws.cell(row=2, column=7).value = excel_serial_1st(year, m)  # G2 统计月份
        if m not in data:
            print("  空白表: %s" % ws.title)
            continue
        raw, daily_map = data[m]
        pos = [r for r in raw if any((s.get('d30') or 0) > 0 for s in r.get('skus', []))]
        fill_month_sheet(ws, pos, daily_map, year, m, store_key, store_dir)
        print("  已填: %s (%d SKC)" % (ws.title, len(pos)))

    # 第一个数据月份 sheet 移到最前，模板 sheet 删除
    first = "%d-%02d" % (year, min(months))
    wb.move_sheet(first, offset=-(wb.sheetnames.index(first)))
    del wb[tpl.title]

    out_name = "%s 每日销量%d.xlsx" % (cfg["name"], year)
    out_path = os.path.join(out_dir, out_name)
    os.makedirs(out_dir, exist_ok=True)
    wb.save(out_path)
    print("已保存:", out_path)
    return out_path

def fill_month_sheet(ws, pos, daily_map, year, month, store_key, store_dir):
    """填一个月份 sheet：数据行 / 删多余行 / SKC+商品名 / 行高 / 产品图"""
    # 当月数据窗口内、属于本月的日期 -> 日列号
    dates = set()
    for skc, mm in daily_map.items():
        for sku in mm.values():
            dates.update(d for d in sku if d.startswith("%d-%02d-" % (year, month)))
    day_col = {}
    for d in dates:
        day = int(d[8:10])
        if 1 <= day <= 31:
            day_col[day] = 11 + day  # L=12 列 = 1日
    if not day_col:
        print("    警告: 本月无任何逐日数据")

    # 清除数据区合并（模板 5~324 行）
    to_unmerge = [str(rng) for rng in ws.merged_cells.ranges if 5 <= rng.min_row and rng.max_row <= 324]
    for r in to_unmerge:
        ws.unmerge_cells(r)

    row = 5
    skc_seq = 0
    for rec in pos:
        skc_id = str(rec['skcId'])
        skus = [s for s in rec.get('skus', []) if (s.get('d30') or 0) > 0]
        if not skus:
            continue
        skc_seq += 1
        first_row = row
        name = rec.get('name', '')
        m = daily_map.get(skc_id, {})
        for s in skus:
            sku_id = str(s['skuId'])
            sku_days = m.get(sku_id, {})
            for day, col in day_col.items():
                v = sku_days.get("%d-%02d-%02d" % (year, month, day), 0)
                if v:
                    ws.cell(row=row, column=col).value = int(v)
            ws.cell(row=row, column=4).value = "SKC: %s\n%s" % (skc_id, name)
            ws.cell(row=row, column=4).alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')
            ws.cell(row=row, column=9).value = s['className']
            ws.cell(row=row, column=11).value = "=SUM(L{r}:AP{r})*J{r}".format(r=row)
            row += 1
        last_row = row - 1
        ws.cell(row=first_row, column=1).value = skc_seq
        ws.merge_cells(start_row=first_row, start_column=2, end_row=last_row, end_column=3)
        ws.merge_cells(start_row=first_row, start_column=4, end_row=last_row, end_column=6)

    last_data_row = row - 1
    if last_data_row < 324:
        ws.delete_rows(last_data_row + 1, 324 - last_data_row)
    residue_start = last_data_row + 5
    if ws.max_row > residue_start:
        ws.delete_rows(residue_start, ws.max_row - residue_start + 1)
    # copy_worksheet 会复制模板全部行高，delete_rows 不会清理 row_dimensions，
    # 需手动删除数据区之后的残留行高（保留表头+总结行），否则保存后残留空行元素
    for r in list(ws.row_dimensions.keys()):
        if r > last_data_row + 4:
            del ws.row_dimensions[r]
    # openpyxl delete_rows 不平移合并单元格：手动重建 merged_cells，
    # 把模板总结行（325-328）的合并平移到新位置（单元格已删，不能 unmerge）
    from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
    offset = 324 - last_data_row
    mcr = MultiCellRange()
    for rng in ws.merged_cells.ranges:
        if rng.min_row >= 325:
            cr = CellRange(str(rng))
            cr.shift(row_shift=-offset, col_shift=0)
            mcr.add(str(cr))
        else:
            mcr.add(str(rng))
    ws.merged_cells = mcr

    for r in range(5, last_data_row + 1):
        ws.row_dimensions[r].height = 64

    # 产品图（当月有销量的 SKC 顺序插入 B:C）
    order = [str(r['skcId']) for r in pos]
    groups = []
    for rng in ws.merged_cells.ranges:
        if rng.min_col == 2 and rng.max_col == 3 and 5 <= rng.min_row <= 71:
            groups.append((rng.min_row, rng.max_row))
    groups.sort()
    if len(groups) != len(order):
        print("    警告: 图片组数 %d != SKC数 %d" % (len(groups), len(order)))
    def col_w_px(w):
        return w * 7 + 5
    region_w = col_w_px(ws.column_dimensions['B'].width or 9.23) + col_w_px(ws.column_dimensions['C'].width or 9.23)
    row_h_px = 64 * 4 / 3
    EMU = 9525
    for skc, (first, last) in zip(order, groups):
        img_path = os.path.join(store_dir, "imgs", "%s_%s.jpg" % (store_key, skc))
        if not os.path.isfile(img_path):
            print("    缺图:", img_path)
            continue
        rows = last - first + 1
        region_h = rows * row_h_px
        img_w = min(region_w - 10, region_h - 8)
        if img_w < 20:
            img_w = 20
        img = XLImage(img_path)
        cx, cy = int(img_w * EMU), int(img_w * EMU)
        dx = int((region_w - img_w) / 2 * EMU)
        dy = int((region_h - img_w) / 2 * EMU)
        img.anchor = OneCellAnchor(_from=AnchorMarker(col=1, colOff=dx, row=first - 1, rowOff=dy),
                                   ext=XDRPositiveSize2D(cx=cx, cy=cy))
        ws.add_image(img)

if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Temu 年度每日销量台账生成器")
    ap.add_argument("--store", choices=list(STORES), default="golf", help="店铺（默认 golf）")
    ap.add_argument("--year", type=int, default=2026)
    ap.add_argument("--template", default=DEFAULT_TPL)
    ap.add_argument("--out-dir", default=DEFAULT_OUT)
    ap.add_argument("--data-dir", default=DEFAULT_DATA)
    args = ap.parse_args()
    build(args.store, args.year, args.template, args.out_dir, args.data_dir)
