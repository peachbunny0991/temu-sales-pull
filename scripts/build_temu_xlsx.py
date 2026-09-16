# -*- coding: utf-8 -*-
"""将浏览器抓取的Temu分SKU销量数据生成为Excel（按近30天销量降序，仅保留有销量数据）。
用法: python3 build_temu_xlsx.py <raw_json> <店铺名> [输出目录]
"""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

raw_file = sys.argv[1]
shop_name = sys.argv[2]
out_dir = sys.argv[3] if len(sys.argv) > 3 else "."

with open(raw_file, encoding="utf-8") as f:
    records = json.load(f)

# ---------- 展平为SKU行（仅保留近30天销量>0） ----------
sku_rows = []
for r in records:
    skc_name = (r.get("name") or "").strip()
    skc_id = r.get("skcId")
    for s in r.get("skus") or []:
        if (s.get("d30") or 0) <= 0:
            continue  # 过滤零销量SKU
        sku_rows.append({
            "skc_name": skc_name, "skc_id": skc_id,
            "sku_name": (s.get("className") or "").strip(),
            "sku_id": s.get("skuId"),
            "today": s.get("today") or 0,
            "d7": s.get("d7") or 0,
            "d30": s.get("d30") or 0,
            "total": s.get("total") or 0,
        })

# 按近30天销量降序，次按累计销量降序
sku_rows.sort(key=lambda x: (-x["d30"], -x["total"]))

# ---------- SKC汇总（仅含有销量SKU的SKC） ----------
skc_map = {}
for r in records:
    skus = [s for s in (r.get("skus") or []) if (s.get("d30") or 0) > 0]
    if not skus:
        continue
    skc_id = r.get("skcId")
    skc_map[skc_id] = {
        "skc_name": (r.get("name") or "").strip(),
        "skc_id": skc_id,
        "sku_cnt": len(skus),
        "d30": sum(s["d30"] for s in skus),
        "total": sum(s["total"] for s in skus),
    }
skc_rows = sorted(skc_map.values(), key=lambda x: (-x["d30"], -x["total"]))

# ---------- 样式 ----------
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
BORDER = Border(*[Side(style="thin", color="D9D9D9")] * 4)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

def write_sheet(ws, headers, rows, widths, center_cols):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT
        cell.alignment = CENTER; cell.border = BORDER
    for ri, row in enumerate(rows, 2):
        for ci, v in enumerate(row, 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.border = BORDER
            cell.alignment = CENTER if ci in center_cols else LEFT
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "A2"

wb = Workbook()
ws1 = wb.active
ws1.title = "分SKU销量明细"
headers1 = ["序号", "SKC名称", "SKC ID", "SKU名称", "SKU ID", "今日销量", "近7天销量", "近30天销量", "累计销量"]
rows1 = [[i + 1, r["skc_name"], r["skc_id"], r["sku_name"], r["sku_id"],
          r["today"], r["d7"], r["d30"], r["total"]] for i, r in enumerate(sku_rows)]
write_sheet(ws1, headers1, rows1, widths=[6, 60, 14, 18, 14, 10, 11, 12, 12], center_cols={1, 3, 4, 5, 6, 7, 8, 9})

ws2 = wb.create_sheet("SKC汇总")
headers2 = ["序号", "SKC名称", "SKC ID", "SKU数", "近30天总销量", "累计总销量"]
rows2 = [[i + 1, r["skc_name"], r["skc_id"], r["sku_cnt"], r["d30"], r["total"]]
         for i, r in enumerate(skc_rows)]
write_sheet(ws2, headers2, rows2, widths=[6, 60, 14, 8, 13, 13], center_cols={1, 3, 4, 5, 6})

out = os.path.join(out_dir, "%s 分SKU近30天销量降序.xlsx" % shop_name)
wb.save(out)
print("saved:", out)
print("有销量SKU行:", len(sku_rows), "| 有销量SKC数:", len(skc_rows))
